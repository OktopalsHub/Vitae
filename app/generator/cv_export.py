"""CV export — DOCX, PDF, and Markdown writers for tailored resumes."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF


def _to_ascii(text: str) -> str:
    return (
        (text or "")
        .replace("\u2014", "-")
        .replace("\u2013", "-")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
    )


# ---------------------------------------------------------------------------
# Markdown export
# ---------------------------------------------------------------------------

def resume_to_markdown(name: str, contact: str, data: dict[str, Any]) -> str:
    lines = [f"# {name}", "", contact, "", "## Professional Summary", "", data.get("summary", ""), "", "## Skills", ""]
    for sk in data.get("skills") or []:
        lines.append(f"- **{sk.get('label', 'Skills')}:** {sk.get('value', '')}")
    lines += ["", "## Professional Experience", ""]
    for exp in data.get("experiences") or []:
        lines.append(f"### {exp.get('title', '')} | {exp.get('dates', '')}")
        lines.append("")
        for b in exp.get("bullets") or []:
            lines.append(f"- {b}")
        lines.append("")
    if data.get("projects"):
        lines += ["## Projects", ""]
        for proj in data["projects"]:
            lines.append(f"### {proj.get('title', '')} | {proj.get('dates', '')}")
            lines.append("")
            for b in proj.get("bullets") or []:
                lines.append(f"- {b}")
            lines.append("")
    lines += ["## Education", ""]
    for edu in data.get("education") or []:
        lines.append(f"**{edu.get('school', '')}** | {edu.get('dates', '')}  ")
        lines.append(edu.get("details", ""))
        lines.append("")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# DOCX writer
# ---------------------------------------------------------------------------

def _add_right_tab(paragraph, pos: str = "9360") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), pos)
    tabs.append(tab)
    p_pr.append(tabs)


def write_docx(path: Path, name: str, contact: str, data: dict[str, Any]) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)

    def job_header(title: str, dates: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10.5)
        p.add_run("\t" + dates)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        _add_right_tab(p)

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = name_p.add_run(name)
    nr.bold = True
    nr.font.size = Pt(18)

    ct = doc.add_paragraph()
    ct.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = ct.add_run(contact)
    cr.font.size = Pt(9.5)

    heading("Professional Summary")
    doc.add_paragraph(data.get("summary", ""))

    heading("Skills")
    for sk in data.get("skills") or []:
        p = doc.add_paragraph()
        r = p.add_run(f"{sk.get('label', 'Skills')}: ")
        r.bold = True
        p.add_run(sk.get("value", ""))

    heading("Professional Experience")
    for exp in data.get("experiences") or []:
        job_header(exp.get("title", ""), exp.get("dates", ""))
        for b in exp.get("bullets") or []:
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)

    if data.get("projects"):
        heading("Projects")
        for proj in data["projects"]:
            job_header(proj.get("title", ""), proj.get("dates", ""))
            for b in proj.get("bullets") or []:
                bp = doc.add_paragraph(b, style="List Bullet")
                bp.paragraph_format.space_after = Pt(1)

    heading("Education")
    for edu in data.get("education") or []:
        job_header(edu.get("school", ""), edu.get("dates", ""))
        doc.add_paragraph(edu.get("details", ""))

    doc.save(str(path))


# ---------------------------------------------------------------------------
# PDF writer
# ---------------------------------------------------------------------------

class ResumePDF(FPDF):
    def section_title(self, title: str) -> None:
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(26, 26, 26)
        self.ln(3)
        self.cell(0, 6, title.upper(), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(40, 40, 40)
        self.set_line_width(0.4)
        y = self.get_y()
        self.line(self.l_margin, y, self.w - self.r_margin, y)
        self.ln(2)

    def job_header(self, title: str, dates: str) -> None:
        self.set_font("Helvetica", "B", 10)
        page_width = self.w - self.l_margin - self.r_margin
        self.cell(page_width * 0.72, 5, _to_ascii(title), align="L")
        self.set_font("Helvetica", "", 9.5)
        self.cell(page_width * 0.28, 5, _to_ascii(dates), align="R", new_x="LMARGIN", new_y="NEXT")

    def bullet(self, text: str) -> None:
        self.set_font("Helvetica", "", 9.5)
        x = self.l_margin
        self.set_x(x)
        self.cell(4, 4.3, "-")
        self.multi_cell(self.w - self.r_margin - x - 4, 4.3, _to_ascii(text))
        self.ln(0.3)


def write_pdf(path: Path, name: str, contact: str, data: dict[str, Any]) -> None:
    pdf = ResumePDF(format="Letter")
    pdf.set_auto_page_break(auto=True, margin=0.45 * 25.4)
    pdf.set_margins(0.6 * 25.4, 0.45 * 25.4, 0.6 * 25.4)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 8, _to_ascii(name), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 5, _to_ascii(contact), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)

    pdf.section_title("Professional Summary")
    pdf.set_font("Helvetica", "", 9.5)
    pdf.multi_cell(0, 4.4, _to_ascii(data.get("summary", "")))

    pdf.section_title("Skills")
    for sk in data.get("skills") or []:
        pdf.set_font("Helvetica", "B", 9.5)
        pdf.write(4.3, f"{sk.get('label', 'Skills')}: ")
        pdf.set_font("Helvetica", "", 9.5)
        pdf.write(4.3, _to_ascii(sk.get("value", "")))
        pdf.ln(4.5)

    pdf.section_title("Professional Experience")
    for exp in data.get("experiences") or []:
        pdf.ln(1)
        pdf.job_header(exp.get("title", ""), exp.get("dates", ""))
        for b in exp.get("bullets") or []:
            pdf.bullet(b)

    if data.get("projects"):
        pdf.section_title("Projects")
        for proj in data["projects"]:
            pdf.ln(1)
            pdf.job_header(proj.get("title", ""), proj.get("dates", ""))
            for b in proj.get("bullets") or []:
                pdf.bullet(b)

    pdf.section_title("Education")
    for edu in data.get("education") or []:
        pdf.ln(1)
        pdf.job_header(edu.get("school", ""), edu.get("dates", ""))
        pdf.set_font("Helvetica", "", 9.5)
        pdf.multi_cell(0, 4.3, _to_ascii(edu.get("details", "")))

    pdf.output(str(path))
